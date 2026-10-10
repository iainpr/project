#!/usr/bin/env python3
"""A local web page for regressions: OLS, panel fixed effects, cross-section, DiD and RDD.

    python portal.py                     # the files in data/, at http://127.0.0.1:8765
    python portal.py --data ~/thesis/data --port 8800 --no-browser

Choose a file, a model and its variables in the page, then run it: summary statistics,
coefficients, diagnostics and charts appear together, and pinned models line up side by
side. The page (portal.html) lists each file's columns and takes the models from MODELS,
so changes happen here, in these sections:

    1. SETTINGS         your column names, variables of interest, controls and defaults
    2. DERIVED          variables made from other columns, such as post: add your own
    3. MODELS           a function per model, and the MODELS table that the page reads
    4. BUILDING BLOCKS  fitting, diagnostics, summary statistics and charts
    5. DATA             reading files, pairing treatment and control data, periods
    6. REQUESTS         checking what the page asks for
    7. SERVER           the local web server (nothing to change)

Everything stays on this computer. The server answers only on 127.0.0.1, checks the Host
header (against DNS rebinding) and a token made for each session (against other
websites), and reads only .xlsx and .csv files directly inside the data folder. Models
are built from arrays, so no column name ever reaches eval or a formula.
"""

from __future__ import annotations

import argparse
import contextlib
import difflib
import functools
import hashlib
import json
import math
import re
import secrets
import threading
import traceback
import warnings
import webbrowser
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from scipy.stats import binomtest
from statsmodels.regression.linear_model import OLS, WLS
from statsmodels.stats.diagnostic import het_breuschpagan, linear_reset
from statsmodels.stats.outliers_influence import variance_inflation_factor
from statsmodels.stats.stattools import durbin_watson, jarque_bera
from statsmodels.tools.tools import add_constant


class UserError(Exception):
    """A problem with the request or the data: the page shows its message."""


# =========================================================================================
# 1. SETTINGS: your data. The page starts from these, and any of them can be changed
#    there. Names match ignoring case, spaces and punctuation, so "CPI Shelter" finds a
#    column called cpi_shelter, and the Data tab lists any name a file lacks. Restart the
#    portal after editing them.
# =========================================================================================

DATA_FOLDER = "data"  # next to this file, or a full path; --data overrides it

UNIT = "census_division"  # one row per unit and period
TIME = "date"  # years, quarters (2015Q1) or dates
TREATMENT_DATE = "treatment_date"  # when each unit's bylaw took effect; blank if never

OUTCOMES = ["permits", "starts"]  # the first one is selected
VARIABLES_OF_INTEREST = ["score"]  # the internal zoning score
CONTROLS = [
    "mortgage_lending",
    "development_charges",
    "bcpi",  # building construction price index
    "nhpi",  # new housing price index
    "cpi_shelter",
    "unemployment_rate",
]

# Treatment and control units can share one sheet (controls have no treatment date), or
# sit apart: in two sheets of one workbook, or in two files. Sheets (and files) whose
# names contain these words are paired, the control data stacked under the treatment data.
TREATMENT_DATA = "treat"  # matches Treatment, Treated, treatment_units.xlsx ...
CONTROL_DATA = "control"

# Where the page starts
MODEL = "ols"  # a key of MODELS: ols, fe, xs, did or rdd
TRANSFORM = "log1p"  # none, log (drops zeros) or log1p: log(1 + y), which keeps zeros
ERRORS = "cluster"  # cluster (by unit), robust or classical, where the model offers it
EVENT_WINDOW = (-8, 8)  # DiD: the periods before and after treatment in the event study
RDD_RUNNING = "score"  # RDD: the running variable
RDD_CUTOFF = None  # RDD: its cutoff, such as 50 (None: enter it in the page)

MAX_MB = 200  # largest file read
POINTS = 2000  # most points drawn in one chart


# =========================================================================================
# 2. DERIVED: variables made from other columns. Each one appears, under its label, in
#    the page's lists of variables of interest and controls. To add one, give it a name, a
#    label and a function of d, where d["column"] is any column as numbers, and
#        d["_unit"]     the unit                                 (with a unit column)
#        d["_t"]        the period, a whole number               (with a time column)
#        d["_adopted"]  the period treatment began, NaN if never (and a treatment date)
# =========================================================================================

DERIVED = {
    "treated": ("Treated unit (has a treatment date)", lambda d: d["_adopted"].notna()),
    "post": ("After treatment (from its period on)", lambda d: d["_t"] >= d["_adopted"]),
    # "score_post": ("Score x after treatment", lambda d: d["score"] * (d["_t"] >= d["_adopted"])),
}


# =========================================================================================
# 3. MODELS: a function per model. Each takes a Job (see section 6: job.data holds the
#    rows, job.xs and job.controls the chosen variables, job.form the model's own inputs)
#    and returns the sections of the results page. To add a model, write a function like
#    these and give it a line in MODELS, at the end of this section.
# =========================================================================================


@dataclass(frozen=True)
class Model:
    label: str  # its button in the page
    fit: Callable[[Job], dict]  # returns title, summary, tables, diagnostics and charts
    hint: str  # the line under the buttons
    inputs: tuple[str, ...] = ("x", "controls")  # the page's input groups (data-input)
    errors: tuple[str, ...] = ("cluster", "robust", "classical")  # standard errors offered
    needs: tuple[str, ...] = ()  # the columns it cannot run without: unit, time, treat


def fit_ols(job: Job) -> dict:
    d = job.sample(job.regressors)
    X = add_constant(d[job.varying(d, job.regressors, "does not vary")], has_constant="add")
    res = fit(d["_y"], X, job.se, d.get("_unit"))
    return dict(
        title=f"OLS: {job.ylabel}",
        summary=[
            ["Observations", res.nobs],
            ["R²", res.rsquared],
            ["Adjusted R²", res.rsquared_adj],
        ],
        tables=[table(f"Coefficients ({ERROR_TITLES[job.se]})", coefficients(res, job.names))],
        diagnostics=diagnostics(res, ordered="_t" in d),
        charts=[partial_chart(job, d["_y"], X, X.columns[1]), residual_chart(res)],
    )


def fit_fe(job: Job) -> dict:
    effects = job.form.options("fe", ("unit", "time"), "Choose unit or time fixed effects, or both")
    if "time" in effects and "_t" not in job.data:
        raise UserError("Time fixed effects need the time column")
    d = job.sample(job.regressors)
    keys = [d["_unit"]] * ("unit" in effects) + [d["_t"]] * ("time" in effects)
    within = demean(d[["_y", *job.regressors]], keys)
    regs = job.varying(within, job.regressors, "is absorbed by the fixed effects")
    res = fit(within["_y"], within[regs], "cluster", d["_unit"])
    fe = " and ".join(e for e in ("unit", "time") if e in effects)
    return dict(
        title=f"Panel fixed effects: {job.ylabel}",
        summary=[
            ["Observations", res.nobs],
            ["Units", d["_unit"].nunique()],
            ["Within R²", within_r2(res, within)],
        ],
        tables=[
            table(
                f"Coefficients ({fe} fixed effects; {ERROR_TITLES['cluster']})",
                coefficients(res, job.names),
            )
        ],
        diagnostics=[
            ["Jarque-Bera p, normal residuals", jarque_bera(res.resid)[1]],
            ["Durbin-Watson, 2 = no autocorrelation", durbin_watson(res.resid)],
        ],
        charts=[partial_chart(job, within["_y"], within[regs], regs[0]), residual_chart(res)],
    )


def fit_xs(job: Job) -> dict:
    regs = job.regressors
    if job.form.choice("xs_mode", ("mean", "period"), "mean") == "mean":
        d, how = job.data.groupby("_unit")[["_y", *regs]].mean(), "unit averages of all periods"
    else:
        if "_t" not in job.data:
            raise UserError("Choose the time column to take one period")
        wanted = pd.Series([job.form.raw.get("xs_period") or ""], name="the period")
        period = time_index(wanted, job.freq)[0].iloc[0]
        d = job.data[job.data["_t"] == period].set_index("_unit")[["_y", *regs]]
        if d.empty:
            raise UserError(f"No rows in period {wanted.iloc[0]!r}")
        how = f"period {period_label(period, job.freq)}"
    d = d.dropna()
    X = add_constant(d[job.varying(d, regs, "does not vary across units")], has_constant="add")
    res = fit(d["_y"], X, job.se)  # one row per unit: robust or classical errors
    return dict(
        title=f"Cross-section: {job.ylabel}",
        summary=[["Units", res.nobs], ["R²", res.rsquared], ["Adjusted R²", res.rsquared_adj]],
        tables=[
            table(
                f"Coefficients, one row per unit: {how} ({ERROR_TITLES[job.se]})",
                coefficients(res, job.names),
            )
        ],
        diagnostics=diagnostics(res, ordered=False),
        charts=[partial_chart(job, d["_y"], X, X.columns[1], d.index), residual_chart(res)],
    )


def fit_did(job: Job) -> dict:
    """Staggered DiD by two-way fixed effects, never-treated units as the comparison: the
    average effect after treatment, and an event study (reference: the period before)."""
    lo = int(job.form.number("lo", "The window's start", EVENT_WINDOW[0]))
    hi = int(job.form.number("hi", "The window's end", EVENT_WINDOW[1]))
    if not lo <= -2 < 0 <= hi:
        raise UserError("The window needs 2 or more periods before treatment and 1 after")
    if built_in := [c for c in job.controls if c in ("treated", "post")]:
        labels = " and ".join(job.names[c] for c in built_in)
        job.notes.append(f"{labels}: part of the DiD itself, so left out of the controls")
    controls = [c for c in job.controls if c not in built_in]
    d = job.data
    early = d["_adopted"] <= d.groupby("_unit")["_t"].transform("min")
    if n := d.loc[early, "_unit"].nunique():
        job.notes.append(f"Units treated in their first period or before, left out: {n}")
    d = d[~early].dropna(subset=["_y", *controls]).copy()
    d.loc[d["_adopted"] > d["_t"].max(), "_adopted"] = np.nan  # treated after the data end
    counts = d.groupby(d["_adopted"].notna())["_unit"].nunique()
    treated, never = counts.get(True, 0), counts.get(False, 0)
    if treated < 2 or never < 2:
        fix = ": choose the control group in step 1" if never < 2 else ""
        raise UserError(
            f"DiD needs 2 or more treated and never-treated units, not {treated} and {never}{fix}"
        )
    since = d["_t"] - d["_adopted"]
    lo, hi = max(lo, int(since.min())), min(hi, int(since.max()))  # no wider than the data
    if lo > -2:
        raise UserError("The treated units need 2 or more periods before treatment")
    event = since.clip(lo, hi)
    d["_post"] = (event >= 0).astype(float)
    ks = [k for k in range(lo, hi + 1) if k != -1 and (event == k).any()]
    if missing := [str(k) for k in range(lo, hi + 1) if k != -1 and k not in ks]:
        job.notes.append(f"No treated rows at t = {', '.join(missing)}, so they are left out")
    events = pd.DataFrame({f"_e{k}": (event == k).astype(float) for k in ks}, index=d.index)
    columns = pd.concat([d[["_y", "_post", *controls]], events], axis=1)
    within = demean(columns, [d["_unit"], d["_t"]])
    controls = job.varying(within, controls, "is absorbed by the fixed effects", need=False)
    res = fit(within["_y"], within[["_post", *controls]], "cluster", d["_unit"])
    es = fit(within["_y"], within[[*events, *controls]], "cluster", d["_unit"])
    leads = [f"_e{k}" for k in ks if k < -1]
    pre_p = wild_p(within[[*events, *controls]], within["_y"], leads, d["_unit"]) if leads else None
    edge = {lo: " or earlier", hi: " or later"}
    names = job.names | {f"_e{k}": f"t = {k}{edge.get(k, '')}" for k in ks}
    rows = coefficients(es, names)[: len(ks)]
    intervals = [[k, r["estimate"], r["low"], r["high"]] for k, r in zip(ks, rows, strict=True)]
    means = d.groupby([d["_adopted"].notna(), "_t"])["_y"].mean()
    trends = []
    for flag, name, color in ((True, "treated units", 1), (False, "never treated", 2)):
        m = means[flag]
        points = np.column_stack([period_x(m.index, job.freq), m])
        labels = [period_label(t, job.freq) for t in m.index]
        trends.append(series(name, "line", points, color, labels=labels))
    job.notes.append(
        "Two-way fixed effects can be biased when effects change over time or differ by "
        "treatment date; project.py has did2s and dCDH, which are not"
    )
    return dict(
        title=f"Difference-in-differences: {job.ylabel}",
        summary=[
            ["Observations", res.nobs],
            ["Treated units", treated],
            ["Never-treated units", never],
            ["Within R²", within_r2(res, within)],
            ["Pre-trend p", pre_p],
        ],
        tables=[
            table(
                f"Average effect after treatment (two-way FE; {ERROR_TITLES['cluster']})",
                coefficients(res, names),
            ),
            table("Event study: the effect by period since treatment (reference t = -1)", rows),
        ],
        diagnostics=[["Pre-trend p: no effects before treatment (wild bootstrap by unit)", pre_p]],
        charts=[
            chart(f"Mean {job.ylabel}, treated and never-treated", "period", job.ylabel, *trends),
            chart(
                "Event study, 95% intervals",
                "periods since treatment",
                "effect",
                series("estimate", "errors", intervals),
                series("reference (t = -1)", "dots", [[-1, 0]], 0),
                hline=0,
                vline=-0.5,
            ),
        ],
    )


def fit_rdd(job: Job) -> dict:
    """Sharp RD: local linear or quadratic fits on each side of the cutoff."""
    running = job.form.column("running", "running variable")
    if (cutoff := job.form.number("cutoff", "The cutoff")) is None:
        raise UserError("Enter the cutoff")
    kernel = job.form.choice("kernel", ("triangular", "uniform"), "triangular")
    order = int(job.form.choice("order", ("1", "2"), "1"))
    controls = [c for c in job.controls if c != running]
    d = job.sample([running, *controls])
    d = d.assign(_r=d[running] - cutoff)
    h = job.form.number("bandwidth", "The bandwidth", float(d["_r"].std()))
    if not h > 0:
        raise UserError("The bandwidth must be positive")
    controls = job.varying(d, controls, "does not vary", need=False)
    args = (kernel, order, controls, job.se)
    res = local_fit(d, h, *args)
    names = job.names | {f"_d{p}": "distance" + "²" * (p - 1) for p in (1, 2)}
    names |= {f"_a{p}": f"{names[f'_d{p}']} x above" for p in (1, 2)}
    fitted = f"Local {('linear', 'quadratic')[order - 1]} fit, {kernel} kernel, bandwidth {h:.4g}"
    tables = [table(f"{fitted} ({ERROR_TITLES[job.se]})", coefficients(res, names))]
    placebo = [
        {"term": job.names[c]} | jump(d.assign(_y=d[c]), h, kernel, order, [], job.se)
        for c in controls
    ]
    if placebo:
        tables.append(table("Placebo: jumps in the controls at the cutoff (expect none)", placebo))
    sensitivity = []
    for factor in (0.5, 0.75, 1, 1.25, 1.5, 2):
        with contextlib.suppress(UserError):  # too few rows at the narrowest bandwidths
            r = jump(d, factor * h, *args)
            sensitivity.append([factor * h, r["estimate"], r["low"], r["high"]])
    # panel rows repeat a unit's running value: count each unit once per value, not each row
    who = "Units" if "_unit" in d else "Rows"
    distinct = d.drop_duplicates(["_unit", "_r"]) if "_unit" in d else d
    near = distinct.loc[distinct["_r"].abs() < h / 4, "_r"]
    density_p = binomtest(int((near >= 0).sum()), len(near)).pvalue if len(near) else None
    inside = distinct[distinct["_r"].abs() < h]

    def count(rows: pd.DataFrame) -> int:
        return rows["_unit"].nunique() if "_unit" in rows else len(rows)

    window = d[d["_r"].abs() <= 2 * h]
    sides = (window[window["_r"] < 0], window[window["_r"] >= 0])
    dots = np.vstack([binned(s["_r"] + cutoff, s["_y"], 15) for s in sides])
    level = res.params["const"] + (res.params[controls] * d[controls].mean()).sum()
    fits = []
    for side in (-1.0, 1.0):
        grid, above = np.linspace(0, side * h, 30), float(side > 0)
        slope = sum(
            (res.params[f"_d{p}"] + above * res.params[f"_a{p}"]) * grid**p
            for p in range(1, order + 1)
        )
        points = np.column_stack([grid + cutoff, level + above * res.params["_jump"] + slope])
        fits.append(series("local fit" if side < 0 else None, "line", points, 2))
    spread = distinct.loc[distinct["_r"].abs() <= 2 * h, "_r"] + cutoff
    counts, edges = np.histogram(spread, bins=30)
    bars = np.column_stack([(edges[:-1] + edges[1:]) / 2, counts])
    label, y = job.names[running], job.ylabel
    return dict(
        title=f"Regression discontinuity: {y} at {label} = {cutoff:g}",
        summary=[
            ["Rows used", res.nobs],
            [f"{who} below the cutoff", count(inside[inside["_r"] < 0])],
            [f"{who} at or above", count(inside[inside["_r"] >= 0])],
            ["Bandwidth", h],
        ],
        tables=tables,
        diagnostics=[
            [f"Density test p, {who.lower()} just below vs above (within h/4)", density_p],
            ["Jarque-Bera p, normal residuals", jarque_bera(res.resid)[1]],
        ],
        charts=[
            chart(
                f"{y} by {label}: bin means and the local fit",
                label,
                y,
                series("bin means", "dots", dots),
                *fits,
                vline=cutoff,
            ),
            chart(
                f"{who} by {label} near the cutoff (bunching check)",
                label,
                who.lower(),
                series(who.lower(), "bars", bars, width=edges[1] - edges[0]),
                vline=cutoff,
            ),
            chart(
                "The jump at other bandwidths, 95% intervals",
                "bandwidth",
                "jump",
                series("jump", "errors", sensitivity),
                hline=0,
                vline=h,
            ),
        ],
    )


def local_fit(d: pd.DataFrame, h: float, kernel: str, order: int, controls: list, se: str) -> Any:
    """Local polynomial regression within h of the cutoff; _jump is the discontinuity."""
    inside = d[d["_r"].abs() < h]
    above = (inside["_r"] >= 0).astype(float)
    if min(above.sum(), len(above) - above.sum()) < order + 3:
        raise UserError(f"Too few rows on one side of the cutoff within bandwidth {h:.4g}")
    X = pd.DataFrame({"_jump": above})
    for p in range(1, order + 1):
        X[f"_d{p}"], X[f"_a{p}"] = inside["_r"] ** p, above * inside["_r"] ** p
    X = add_constant(X.join(inside[controls]), has_constant="add")
    weights = 1 - inside["_r"].abs() / h if kernel == "triangular" else None
    return fit(inside["_y"], X, se, inside.get("_unit"), weights)


def jump(d: pd.DataFrame, h: float, kernel: str, order: int, controls: list, se: str) -> dict:
    """The jump at the cutoff, as a coefficient row."""
    return coefficients(local_fit(d, h, kernel, order, controls, se), {})[1]


MODELS = {  # in the page's order; the key is what the page sends
    "ols": Model(
        "OLS",
        fit_ols,
        "Pooled OLS of the outcome on the variables of interest and controls.",
    ),
    "fe": Model(
        "Panel FE",
        fit_fe,
        "The within estimator, with unit fixed effects, time fixed effects or both.",
        inputs=("x", "controls", "fe"),
        errors=("cluster",),
        needs=("unit",),
    ),
    "xs": Model(
        "Cross-section",
        fit_xs,
        "One row per unit: unit averages (the between estimator) or a single period.",
        inputs=("x", "controls", "xs"),
        errors=("robust", "classical"),
        needs=("unit",),
    ),
    "did": Model(
        "DiD",
        fit_did,
        "Staggered difference-in-differences against never-treated units: the average "
        "effect, an event study and a pre-trend test.",
        inputs=("controls", "window"),
        errors=("cluster",),
        needs=("unit", "time", "treat"),
    ),
    "rdd": Model(
        "RDD",
        fit_rdd,
        "Sharp regression discontinuity: local linear or quadratic fits on each side of a cutoff.",
        inputs=("controls", "rdd"),
    ),
}


# =========================================================================================
# 4. BUILDING BLOCKS: fitting, diagnostics, summary statistics and charts.
# =========================================================================================

TRANSFORMS = {"none": "{}", "log": "log({})", "log1p": "log(1 + {})"}
ERROR_CHOICES = {"cluster": "clustered by unit", "robust": "robust (HC1)", "classical": "classical"}
ERROR_TITLES = {
    "cluster": "errors clustered by unit",
    "robust": "robust errors",
    "classical": "classical errors",
}
TERMS = {"const": "Constant", "_post": "Post (treated)", "_jump": "Jump at the cutoff"}


def fit(y: pd.Series, X: pd.DataFrame, se: str, groups: Any = None, weights: Any = None) -> Any:
    """OLS (WLS with weights) with errors clustered by unit, robust (HC1) or classical.
    Clustered tests use t and F with units - 1 degrees of freedom, as Stata does."""
    model = OLS(y, X) if weights is None else WLS(y, X, weights=weights)
    if se == "cluster":
        groups = pd.factorize(groups)[0]
        return model.fit(cov_type="cluster", cov_kwds={"groups": groups}, use_t=True)
    return model.fit(cov_type="HC1" if se == "robust" else "nonrobust", use_t=True)


def wild_p(X: pd.DataFrame, y: pd.Series, tested: list[str], groups: Any, reps: int = 999) -> float:
    """p-value for "the tested coefficients are all zero" from a wild score bootstrap that
    flips the sign of each unit's score (Kline and Santos 2012). The clustered F test
    rejects far too often when few units are treated; this one keeps its size."""
    rest = X.drop(columns=tested).to_numpy()

    def residual(v: np.ndarray) -> np.ndarray:
        return v - rest @ np.linalg.lstsq(rest, v, rcond=None)[0]

    u, Z = residual(y.to_numpy()), residual(X[tested].to_numpy())
    codes = pd.factorize(groups)[0]
    H = np.zeros((codes.max() + 1, len(tested)))
    np.add.at(H, codes, Z * u[:, None])  # each unit's score, under the null
    A = H @ np.linalg.pinv(H.T @ H) @ H.T
    signs = np.random.default_rng(0).choice([-1.0, 1.0], (reps, len(H)))  # same p each run
    draws = np.einsum("rg,gh,rh->r", signs, A, signs)
    return float((1 + (draws >= A.sum() * (1 - 1e-9)).sum()) / (reps + 1))


def demean(frame: pd.DataFrame, keys: list[pd.Series], tol: float = 1e-9) -> pd.DataFrame:
    """Remove fixed effects by subtracting group means for each key in turn until stable:
    one round for one key or a balanced panel, a few more for an unbalanced one."""
    out = frame.astype(float)
    scale = float(out.abs().max().max()) + 1.0
    for _ in range(1000):
        before = out
        for key in keys:
            out = out - out.groupby(key.to_numpy()).transform("mean")
        if len(keys) == 1 or float((out - before).abs().max().max()) < tol * scale:
            return out
    warnings.warn("The fixed effects did not fully converge", stacklevel=2)
    return out


def within_r2(res: Any, within: pd.DataFrame) -> float:
    return 1 - res.ssr / float((within["_y"] ** 2).sum())


def table(title: str, rows: list[dict]) -> dict:
    return {"title": title, "rows": rows}


def coefficients(res: Any, names: dict[str, str]) -> list[dict]:
    """A coefficient table's rows: estimate, standard error, t, p and the 95% interval."""
    ci = res.conf_int()
    stats = {"estimate": res.params, "se": res.bse, "t": res.tvalues, "p": res.pvalues}
    frame = pd.DataFrame(stats | {"low": ci[0], "high": ci[1]})
    return [{"term": names.get(t, t)} | row for t, row in frame.to_dict("index").items()]


def diagnostics(res: Any, ordered: bool) -> list[list]:
    """Checks of an OLS fit with a constant; add rows here to show more."""
    exog = res.model.exog
    rows = [
        ["Breusch-Pagan p, constant variance", het_breuschpagan(res.resid, exog)[1]],
        ["Jarque-Bera p, normal residuals", jarque_bera(res.resid)[1]],
        ["RESET p, functional form", linear_reset(res, power=2, use_f=True).pvalue],
    ]
    if exog.shape[1] > 2:
        vif = max(variance_inflation_factor(exog, i) for i in range(1, exog.shape[1]))
        rows.append(["Largest variance inflation factor", vif])
    if ordered:
        rows.append(["Durbin-Watson, 2 = no autocorrelation", durbin_watson(res.resid)])
    return rows


def summary_stats(job: Job) -> dict:
    """Count, mean, SD, min and max of the outcome and each variable in the model. With a
    treatment date, also the treated and control means before the first treatment, and
    their standardized difference."""
    columns = ["_y", *job.described]
    d = job.data.dropna(subset=columns)
    header = ["Variable", "N", "Mean", "SD", "Min", "Max"]
    groups = None
    if "_adopted" in d and d["_adopted"].notna().any() and d["_adopted"].isna().any():
        first = d["_adopted"].min()
        if len(before := d[d["_t"] < first]):
            groups = before[before["_adopted"].notna()], before[before["_adopted"].isna()]
            when = period_label(first, job.freq)
            header += [f"Treated, before {when}", "Control, before", "Std. diff."]
    rows = []
    for c in columns:
        x = d[c]
        name = job.ylabel if c == "_y" else job.names[c]
        row = [name, x.count(), x.mean(), x.std(), x.min(), x.max()]
        if groups:
            treated, control = groups[0][c], groups[1][c]
            spread = math.sqrt((treated.var() + control.var()) / 2)
            gap = (treated.mean() - control.mean()) / spread if spread > 0 else None
            row += [treated.mean(), control.mean(), gap]
        rows.append(row)
    note = (
        "Std. diff.: the difference in means over their pooled SD; beyond about 0.25 either "
        "way, the groups differ noticeably."
    )
    return {"title": "Summary statistics", "columns": header, "rows": rows} | {
        "note": note if groups else None
    }


def series(name: str | None, kind: str, points: Any, color: int = 1, **extra: Any) -> dict:
    """A chart's series. kind: dots, line, errors ([x, y, low, high] points) or bars."""
    return {
        "name": name,
        "kind": kind,
        "points": np.asarray(points).tolist(),
        "color": color,
    } | extra


def chart(title: str, x: str, y: str, *lines: dict, **reference: float) -> dict:
    """A chart for the page to draw; hline= and vline= add reference lines."""
    return {"title": title, "x": x, "y": y, "series": list(lines)} | reference


def binned(x: pd.Series, y: pd.Series, bins: int = 20) -> np.ndarray:
    data = pd.DataFrame({"x": np.asarray(x), "y": np.asarray(y)}).dropna()
    few = data["x"].nunique() <= bins  # then a bin per value (the values, not the column)
    key = data["x"].to_numpy() if few else pd.qcut(data["x"], bins, duplicates="drop")
    return data.groupby(key, observed=True).mean()[["x", "y"]].to_numpy()


def partial_chart(job: Job, y: pd.Series, X: pd.DataFrame, focus: str, labels: Any = None) -> dict:
    """y against focus with the other regressors partialled out, so the line's slope is
    focus's coefficient (Frisch-Waugh-Lovell): bin means, or each unit when few."""
    others, x = X.drop(columns=focus), X[focus]
    if others.shape[1]:
        y = y - OLS(y, others).fit().fittedvalues + y.mean()
        x = x - OLS(x, others).fit().fittedvalues + x.mean()
    slope, intercept = np.polyfit(x, y, 1)
    ends = np.array([x.min(), x.max()])
    line = series("fitted line", "line", np.column_stack([ends, intercept + slope * ends]), 2)
    if labels is not None and len(x) <= POINTS:
        dots = series("units", "dots", np.column_stack([x, y]), labels=[str(u) for u in labels])
    else:
        dots = series("bin means", "dots", binned(x, y))
    name, held = (
        job.names[focus],
        ", other regressors held fixed" if set(others) - {"const"} else "",
    )
    return chart(f"{job.ylabel} against {name}{held}", name, job.ylabel, dots, line)


def residual_chart(res: Any) -> dict:
    # at most POINTS rows, the same ones each run
    keep = np.sort(np.random.default_rng(0).permutation(len(res.resid))[:POINTS])
    points = np.column_stack([res.fittedvalues.to_numpy()[keep], res.resid.to_numpy()[keep]])
    dots = series("rows", "dots", points, small=True)
    return chart("Residuals against fitted values", "fitted value", "residual", dots, hline=0)


# =========================================================================================
# 5. DATA: reading files, pairing treatment and control data, periods, column summaries.
# =========================================================================================

HINTS = {  # when a setting is not in a file: column-name fragments that suggest a role
    "unit": ("census_division", "division", "municipal", "geo", "region", "unit", "cd"),
    "time": ("ref_date", "date", "period", "quarter", "year", "month", "time"),
    "treat": ("treat", "effective", "adopt", "bylaw"),
    "y": ("permit", "start"),
}


def clean(name: object) -> str:
    return re.sub(r"[^0-9a-z]+", "_", str(name).lower()).strip("_") or "column"


def unique(names: list[str]) -> Iterator[str]:
    """Number repeated names: a, a -> a, a_2."""
    seen: dict[str, int] = {}
    for name in names:
        seen[name] = seen.get(name, 0) + 1
        yield name if seen[name] == 1 else f"{name}_{seen[name]}"


def data_files(folder: Path) -> dict[str, Path]:
    """The .xlsx and .csv files directly in folder, without hidden or Excel lock files."""
    paths = sorted(folder.iterdir(), key=lambda p: p.name.lower())
    ok = (".xlsx", ".csv")
    return {p.name: p for p in paths if p.suffix.lower() in ok and p.name[0] not in ".~"}


@functools.lru_cache(maxsize=16)
def sheets(path: Path, mtime: float) -> list[str]:
    """A workbook's sheet names ([] for a CSV file)."""
    if path.suffix.lower() != ".xlsx":
        return []
    from openpyxl import load_workbook  # parses the XML with defusedxml when installed

    book = load_workbook(path, read_only=True)
    try:
        return list(book.sheetnames)
    finally:
        book.close()


@functools.lru_cache(maxsize=4)
def read(path: Path, sheet: str, mtime: float) -> tuple[pd.DataFrame, tuple[str, ...]]:
    """A sheet or CSV file with clean, unique column names, cached until the file changes.

    Requests share the cached frame, so nothing may change it in place.
    """
    if path.stat().st_size > MAX_MB * 2**20:
        raise UserError(f"{path.name} is larger than {MAX_MB} MB")
    if path.suffix.lower() == ".xlsx":
        frame = pd.read_excel(path, sheet_name=sheet)
    else:
        try:
            frame = pd.read_csv(path, thousands=",", encoding="utf-8-sig")
        except UnicodeDecodeError:  # a CSV saved by Excel on Windows
            frame = pd.read_csv(path, thousands=",", encoding="cp1252")
    labels = tuple(str(c) for c in frame.columns)
    frame.columns = list(unique([clean(c) for c in labels]))
    return frame.dropna(how="all").reset_index(drop=True), labels


def combine(treated: tuple, control: tuple) -> tuple[pd.DataFrame, tuple[str, ...], list[str]]:
    """The control data stacked under the treatment data, columns matched by clean name;
    _group says which each row came from."""
    (a, a_labels), (b, b_labels) = treated, control
    label = dict(zip(b.columns, b_labels, strict=True))
    label |= dict(zip(a.columns, a_labels, strict=True))  # the treatment data's names win
    only = {"treatment": [c for c in a if c not in b], "control": [c for c in b if c not in a]}
    notes = [
        f"Only in the {group} data: {', '.join(label[c] for c in columns)}"
        for group, columns in only.items()
        if columns
    ]
    frame = pd.concat([a.assign(_group="treatment"), b.assign(_group="control")], ignore_index=True)
    return frame, tuple(label.get(c, c) for c in frame.columns), notes


def pairing(books: dict[str, list[str]]) -> tuple[str | None, dict[str, dict]]:
    """The file to open first, and for each file the sheet to open and where its control
    data are, matching TREATMENT_DATA and CONTROL_DATA in sheet and file names: a control
    sheet in the same workbook, or else (for a treatment file) a control file."""

    def has(name: str, word: str) -> bool:
        return bool(word) and clean(word) in clean(name)

    def pick(book: list[str], word: str) -> str:
        return next((s for s in book if has(s, word)), book[0] if book else "")

    pairs = {}
    for name, book in books.items():
        sheet = pick(book, TREATMENT_DATA)
        control = next(
            ({"file": name, "sheet": s} for s in book if s != sheet and has(s, CONTROL_DATA)), None
        )
        if control is None and has(Path(name).stem, TREATMENT_DATA):
            other = next((n for n in books if n != name and has(Path(n).stem, CONTROL_DATA)), None)
            control = {"file": other, "sheet": pick(books[other], CONTROL_DATA)} if other else None
        pairs[name] = {"sheet": sheet, "control": control}
    treatment = [
        name
        for name, book in books.items()
        if has(Path(name).stem, TREATMENT_DATA) or any(has(s, TREATMENT_DATA) for s in book)
    ]
    return next(iter(treatment or books), None), pairs


def to_number(values: pd.Series) -> pd.Series:
    """Numbers as floats; text such as StatCan's ".." or "x" becomes missing."""
    if pd.api.types.is_numeric_dtype(values):
        return values.astype(float)
    text = values.astype("string").str.replace(",", "", regex=False)
    return pd.to_numeric(text, errors="coerce").astype(float)


def to_dates(values: pd.Series) -> pd.Series:
    """Years, quarters (2015Q1, Q1 2015) and dates as timestamps; NaT where unreadable."""
    text = values.astype("string").str.strip().str.upper().str.removesuffix(".0")
    text = text.str.replace(r"^Q([1-4])\s*-?\s*(\d{4})$", r"\2Q\1", regex=True)
    text = text.str.replace(r"^(\d{4})\s*-?\s*Q([1-4])$", r"\1Q\2", regex=True)
    return pd.to_datetime(text.replace("", pd.NA), errors="coerce", format="mixed")


def kind(values: pd.Series) -> str:
    """number, date or text, judged on a sample of the values."""
    if pd.api.types.is_numeric_dtype(values):
        return "number"
    if pd.api.types.is_datetime64_any_dtype(values):
        return "date"
    sample = values.dropna().head(200)
    if sample.empty:
        return "text"
    if to_number(sample).notna().mean() > 0.9:
        return "number"
    return "date" if to_dates(sample).notna().mean() > 0.9 else "text"


def time_index(values: pd.Series, freq: str | None = None) -> tuple[pd.Series, str]:
    """Periods as integer ordinals, and the frequency: Y, Q, M or # (plain integers).

    The frequency follows from the spacing of the dates unless it is given, as it is for
    treatment dates (the time column's frequency). Blank cells are missing.
    """
    filled = values.astype("string").str.strip().fillna("") != ""
    if freq == "#" or (freq is None and pd.api.types.is_numeric_dtype(values)):
        numbers = to_number(values)
        if (filled & numbers.isna()).any() or not (numbers.dropna() % 1 == 0).all():
            raise UserError(f"{values.name}: periods must be whole numbers or dates")
        if freq == "#" or not numbers.dropna().between(1000, 9999).all():
            return numbers, "#"
    dates = to_dates(values)
    if (unread := filled & dates.isna()).any():
        raise UserError(f"Cannot read {values[unread].iloc[0]!r} in {values.name} as a date")
    if freq is None:
        months = np.unique((dates.dt.year * 12 + dates.dt.month).dropna().to_numpy())
        if len(months) < dates.dropna().nunique():
            raise UserError(f"{values.name} has several dates in a month: use months or longer")
        step = int(np.gcd.reduce(np.diff(months).astype(int))) if len(months) > 1 else 12
        freq = "Y" if step % 12 == 0 else "Q" if step % 3 == 0 else "M"
    ordinals = dates.dt.to_period(freq).array.asi8.astype(float)
    ordinals[dates.isna().to_numpy()] = np.nan
    return pd.Series(ordinals, index=values.index), freq


def period_label(ordinal: float, freq: str) -> str:
    return f"{ordinal:g}" if freq == "#" else str(pd.Period(ordinal=int(ordinal), freq=freq))


def period_x(ordinals: Any, freq: str) -> np.ndarray:
    """Periods as decimal years for charts: 2015Q2 -> 2015.25."""
    if freq == "#":
        return np.asarray(ordinals, dtype=float)
    start = pd.PeriodIndex.from_ordinals(np.asarray(ordinals, dtype=int), freq=freq).start_time
    return np.asarray(start.year + (start.month - 1) / 12, dtype=float)


def describe(frame: pd.DataFrame, labels: tuple[str, ...]) -> dict:
    """Each column's kind and summary (the Data tab), and the rows from each group when
    the treatment and control data come from separate sheets or files."""
    columns = []
    for name, label in zip(frame.columns, labels, strict=True):
        if name.startswith("_"):
            continue
        values = frame[name]
        entry = {"name": name, "label": label, "kind": kind(values)}
        entry |= {"missing": int(values.isna().sum()), "unique": int(values.nunique())}
        if entry["kind"] == "number":
            x = to_number(values)
            entry |= {"mean": x.mean(), "sd": x.std(), "min": x.min(), "max": x.max()}
        columns.append(entry)
    groups = frame["_group"].value_counts().to_dict() if "_group" in frame else {}
    return {"rows": len(frame), "columns": columns, "groups": groups}


def guess(role: str, columns: list[dict]) -> str | None:
    """A column whose name suggests the role (HINTS), for files the settings do not fit."""
    kinds = {"unit": "text number", "time": "date number text", "treat": "date text number"}
    for word in HINTS[role]:
        for c in columns:
            avoid = role == "time" and any(w in c["name"] for w in HINTS["treat"])
            if word in c["name"] and c["kind"] in kinds.get(role, "number") and not avoid:
                return c["name"]
    return None


def settings_for(columns: list[dict]) -> tuple[dict, list[dict]]:
    """The SETTINGS matched to a table's columns: where the page starts, and a check of
    each name (found, or why not). Roles the settings miss are guessed from HINTS."""
    label = {c["name"]: c["label"] for c in columns}
    numbers = [c["name"] for c in columns if c["kind"] == "number"]
    check = []

    def find(setting: str, name: str, pool: list[str]) -> str | None:
        key = clean(name)
        if found := key in pool:
            note = None
        elif key in label:
            note = "is not numbers"
        else:
            close = difflib.get_close_matches(key, list(label), n=1, cutoff=0.6)
            note = f"closest: {label[close[0]]}" if close else None
        check.append({"setting": setting, "name": name, "found": label.get(key) if found else None})
        check[-1]["note"] = note
        return key if found else None

    unit = find("UNIT", UNIT, list(label)) or guess("unit", columns)
    time = find("TIME", TIME, list(label)) or guess("time", columns)
    treat = find("TREATMENT_DATE", TREATMENT_DATE, list(label)) or guess("treat", columns)
    outcomes = [c for n in OUTCOMES if (c := find("OUTCOMES", n, numbers))]
    xs = [c for n in VARIABLES_OF_INTEREST if (c := find("VARIABLES_OF_INTEREST", n, numbers))]
    controls = [c for n in CONTROLS if (c := find("CONTROLS", n, numbers))]
    running = find("RDD_RUNNING", RDD_RUNNING, numbers) if RDD_RUNNING else None
    start = {
        "model": MODEL,
        "transform": TRANSFORM,
        "se": ERRORS,
        "unit": unit,
        "time": time,
        "treat": treat,
        "y": outcomes[0] if outcomes else guess("y", columns),
        "x": xs,
        "controls": controls,
        "running": running or (xs[0] if xs else None),
        "cutoff": RDD_CUTOFF,
        "lo": EVENT_WINDOW[0],
        "hi": EVENT_WINDOW[1],
    }
    return start, check


# =========================================================================================
# 6. REQUESTS: the page's request, checked against the table, becomes a Job.
# =========================================================================================

ROLES = (("unit", "unit column"), ("time", "time column"), ("treat", "treatment date column"))
NEEDS = {"_unit": "the unit column", "_t": "the time column", "_adopted": "the treatment date"}


class Form:
    """The page's request, checked field by field; each field is an input in portal.html."""

    def __init__(self, raw: dict, columns: list[str], derived: list[str]) -> None:
        self.raw = raw
        self.columns = {c for c in columns if not c.startswith("_")}
        self.allowed = self.columns | set(derived)

    def column(self, key: str, what: str, required: bool = True) -> str | None:
        """A column of the table (a menu in the page), or None if left blank."""
        name = self.raw.get(key) or None
        if name is None and required:
            raise UserError(f"Choose the {what}")
        if name is not None and not (isinstance(name, str) and name in self.columns):
            raise UserError(f"Unknown column for the {what}: {name!r}")
        return name

    def variables(self, key: str) -> list[str]:
        """Columns or DERIVED variables (a checklist in the page)."""
        names = self.raw.get(key) or []
        if not isinstance(names, list) or not all(isinstance(n, str) for n in names):
            raise UserError(f"{key} must be a list of column names")
        if unknown := [n for n in names if n not in self.allowed]:
            raise UserError(f"Unknown column in {key}: {unknown[0]!r}")
        return list(dict.fromkeys(names))

    def choice(self, key: str, options: Any, default: str) -> str:
        value = str(self.raw.get(key) or default)
        if value not in options:
            raise UserError(f"{key} must be one of {', '.join(options)}")
        return value

    def options(self, key: str, allowed: tuple[str, ...], message: str) -> list[str]:
        """One or more of allowed (checkboxes in the page)."""
        values = self.raw.get(key, list(allowed))
        if not isinstance(values, list) or not values or not {*map(str, values)} <= set(allowed):
            raise UserError(message)
        return values

    def number(self, key: str, what: str, default: float | None = None) -> float | None:
        value = self.raw.get(key)
        if value in (None, ""):
            return default
        try:
            number = float(value)
        except (TypeError, ValueError):
            raise UserError(f"{what} must be a number") from None
        if not math.isfinite(number):
            raise UserError(f"{what} must be a finite number")
        return number


@dataclass
class Job:
    """A checked request: the analysis table, and what the page asked for."""

    model: Model
    form: Form  # the model's own inputs, such as job.form.number("cutoff", ...)
    frame: pd.DataFrame  # the table as read (shared between requests, so read-only)
    data: pd.DataFrame  # _y, the chosen variables, and _unit, _t and _adopted if chosen
    xs: list[str]  # variables of interest
    controls: list[str]
    described: list[str]  # the variables in the summary statistics
    names: dict[str, str]  # column -> label shown
    ylabel: str
    se: str
    freq: str | None
    notes: list[str]

    @property
    def regressors(self) -> list[str]:
        return self.xs + self.controls

    def sample(self, columns: list[str]) -> pd.DataFrame:
        """The rows that have the outcome and every one of columns."""
        return self.data.dropna(subset=["_y", *columns])

    def varying(self, frame: pd.DataFrame, columns: list[str], why: str, need: bool = True) -> list:
        """The columns that vary in frame; the others are noted and left out."""
        keep = [c for c in columns if frame[c].std() > 1e-9 * (1 + frame[c].abs().max())]
        self.notes += [
            f"{self.names[c]} {why}, so it is left out" for c in columns if c not in keep
        ]
        if need and not keep:
            raise UserError(f"Nothing left to estimate: every regressor {why}")
        return keep


class Columns:
    """What a DERIVED function sees: d[name] is a column as numbers (or _unit, _t, _adopted)."""

    def __init__(self, data: pd.DataFrame, frame: pd.DataFrame) -> None:
        self.data, self.frame = data, frame

    def __getitem__(self, name: str) -> pd.Series:
        if name in self.data:
            return self.data[name]
        if name in self.frame and not name.startswith("_"):
            return to_number(self.frame[name]).loc[self.data.index]
        raise KeyError(name)


def derive(name: str, data: pd.DataFrame, frame: pd.DataFrame) -> pd.Series:
    label, make = DERIVED[name]
    try:
        values = make(Columns(data, frame))
    except KeyError as exc:
        need = exc.args[0]
        raise UserError(f"{label} needs {NEEDS.get(need, f'a column {need!r}')}") from None
    return pd.Series(values, index=data.index).astype(float)


def check_groups(group: pd.Series, units: pd.Series) -> None:
    """A unit belongs to the treatment data or the control data, not both."""
    if both := set(units[group == "treatment"]) & set(units[group == "control"]):
        raise UserError(
            f"{sorted(both)[0]} is in both the treatment and the control data: "
            "units need distinct names"
        )


def group_notes(group: pd.Series, data: pd.DataFrame) -> list[str]:
    """Treatment units without a treatment date, and control units with one."""
    undated = data.loc[(group == "treatment") & data["_adopted"].isna(), "_unit"].nunique()
    dated = data.loc[(group == "control") & data["_adopted"].notna(), "_unit"].nunique()
    notes = [f"Treatment units without a treatment date, so counted as controls: {undated}"]
    notes += [f"Control units with a treatment date, so counted as treated: {dated}"]
    return [note for note, n in zip(notes, (undated, dated), strict=True) if n]


def make_job(frame: pd.DataFrame, labels: tuple[str, ...], raw: dict) -> Job:
    """Check the page's request against the table, and build the analysis table."""
    derived = [k for k in DERIVED if k not in frame]  # a column of the same name wins
    form = Form(raw, list(frame.columns), derived)
    model = MODELS[form.choice("model", MODELS, MODEL)]
    names = {k: DERIVED[k][0] for k in derived} | dict(zip(frame.columns, labels, strict=True))
    names |= TERMS
    y = form.column("y", "outcome")
    xs = [c for c in form.variables("x") if c != y] if "x" in model.inputs else []
    controls = [c for c in form.variables("controls") if c not in (y, *xs)]
    controls = controls if "controls" in model.inputs else []
    if "x" in model.inputs and not xs + controls:
        raise UserError("Choose at least one variable of interest or control")
    running = form.column("running", "running variable") if "rdd" in model.inputs else None
    se = form.choice("se", model.errors, model.errors[0])
    unit, time, treat = (form.column(role, what, role in model.needs) for role, what in ROLES)
    if se == "cluster" and not unit:
        raise UserError("Clustered errors need the unit column: choose it, or robust errors")
    transform = form.choice("transform", TRANSFORMS, "none")
    notes: list[str] = []
    data = pd.DataFrame(index=frame.index)
    for column in dict.fromkeys([y, *xs, *controls, *([running] if running else [])]):
        if column in derived:
            continue  # made below, once the panel columns are in place
        data[column] = to_number(frame[column])
        if bad := int((data[column].isna() & frame[column].notna()).sum()):
            notes.append(f"{names[column]}: {bad} values that are not numbers count as missing")
    data["_y"] = data[y]
    if transform != "none":
        invalid = data[y] <= 0 if transform == "log" else data[y] < 0
        if invalid.any():
            notes.append(f"The {transform} transform leaves out {int(invalid.sum())} rows")
        data["_y"] = (np.log if transform == "log" else np.log1p)(data[y].where(~invalid))
    freq = None
    if unit:
        data["_unit"] = frame[unit].astype("string").str.strip().str.removesuffix(".0")
        if "_group" in frame:
            check_groups(frame["_group"], data["_unit"])
    if time:
        data["_t"], freq = time_index(frame[time].rename(names[time]))
    if keys := [c for c in ("_unit", "_t") if c in data]:
        data = data.dropna(subset=keys).sort_values(keys)  # in time order within each unit
    if unit and time and data.duplicated(keys).any():
        raise UserError(f"Some units repeat a period of {names[time]}: one row per unit and period")
    if treat and not time:
        notes.append("The treatment date is used only with a time column")
    elif treat:
        adopted = time_index(frame[treat].rename(names[treat]), freq)[0].loc[data.index]
        if unit:
            if (adopted.groupby(data["_unit"]).nunique() > 1).any():
                raise UserError(f"Some units have more than one {names[treat]}")
            adopted = adopted.groupby(data["_unit"]).transform("first")
        data["_adopted"] = adopted
    if "_group" in frame and "_unit" in data and "_adopted" in data:
        notes += group_notes(frame["_group"].loc[data.index], data)
    for column in dict.fromkeys([*xs, *controls]):
        if column in derived:
            data[column] = derive(column, data, frame)
    described = list(dict.fromkeys([*xs, *controls, *([running] if running else [])]))
    ylabel = TRANSFORMS[transform].format(names[y])
    return Job(model, form, frame, data, xs, controls, described, names, ylabel, se, freq, notes)


# =========================================================================================
# 7. SERVER: the local web server. Nothing here needs changing for new data or models.
# =========================================================================================

PAGE = Path(__file__).with_name("portal.html")
MAX_BODY = 64 * 1024  # largest request accepted
SETTINGS = [UNIT, TIME, TREATMENT_DATE, OUTCOMES, VARIABLES_OF_INTEREST, CONTROLS, CONTROL_DATA]
SETTINGS += [TREATMENT_DATA, MODEL, TRANSFORM, ERRORS, EVENT_WINDOW, RDD_RUNNING, RDD_CUTOFF]
SETTINGS_ID = hashlib.sha256(repr([*SETTINGS, *DERIVED]).encode()).hexdigest()[:12]


def jsonable(value: Any) -> Any:
    """Plain JSON: floats to 6 significant digits, NaN and infinity as null."""
    if isinstance(value, float | np.floating):
        return float(f"{value:.6g}") if math.isfinite(value) else None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, dict):
        return {str(k): jsonable(v) for k, v in value.items()}
    if isinstance(value, list | tuple):
        return [jsonable(v) for v in value]
    return value


class Portal(ThreadingHTTPServer):
    daemon_threads = True

    def __init__(self, folder: Path, port: int) -> None:
        super().__init__(("127.0.0.1", port), Handler)
        self.folder, self.token, self.lock = folder, secrets.token_urlsafe(32), threading.Lock()
        self.port = self.server_address[1]
        self.hosts = {f"127.0.0.1:{self.port}", f"localhost:{self.port}"}

    def table(self, files: dict[str, Path], name: Any, sheet: Any) -> tuple:
        """A listed file's sheet (or CSV file) as read: frame, labels, sheet, sheet names."""
        if not isinstance(name, str) or name not in files:  # a listed name only: no paths
            raise UserError("Choose a file from the list")
        path = files[name]
        book = sheets(path, path.stat().st_mtime)
        sheet = (sheet or book[0]) if book else ""
        if sheet not in [*book, ""]:
            raise UserError(f"No sheet {sheet!r} in {name}")
        try:
            frame, labels = read(path, sheet, path.stat().st_mtime)
        except UserError:
            raise
        except Exception as exc:  # pandas and openpyxl raise many types for unreadable files
            raise UserError(f"Cannot read {name}: {exc}") from None
        return frame, labels, sheet, book

    def load(self, form: dict) -> tuple[pd.DataFrame, tuple[str, ...], dict]:
        """The table the page chose, with the control data stacked under it when they are
        in another sheet or file."""
        files = data_files(self.folder)
        frame, labels, sheet, book = self.table(files, form.get("file"), form.get("sheet"))
        about = {"file": form["file"], "sheets": book, "sheet": sheet, "notes": []}
        if form.get("control_file"):
            control = self.table(files, form["control_file"], form.get("control_sheet"))
            if (form["control_file"], control[2]) == (form["file"], sheet):
                raise UserError("The control data must be another sheet or file")
            frame, labels, about["notes"] = combine((frame, labels), control[:2])
        return frame, labels, about

    def files(self, form: dict) -> dict:
        found = data_files(self.folder)
        books = {}
        for name, path in found.items():
            try:
                books[name] = sheets(path, path.stat().st_mtime)
            except Exception:  # an unreadable workbook: the error shows when it is opened
                books[name] = []
        first, pairs = pairing(books)
        models = {
            key: {"label": m.label, "hint": m.hint, "inputs": m.inputs, "errors": m.errors}
            for key, m in MODELS.items()
        }
        return {
            "folder": str(self.folder),
            "files": [
                {"name": n, "bytes": p.stat().st_size, "sheets": books[n]} | pairs[n]
                for n, p in found.items()
            ],
            "open": first,
            "models": models,
            "errors": ERROR_CHOICES,
            "derived": {name: label for name, (label, _) in DERIVED.items()},
            "settings_id": SETTINGS_ID,
        }

    def profile(self, form: dict) -> dict:
        frame, labels, about = self.load(form)
        summary = describe(frame, labels)
        defaults, check = settings_for(summary["columns"])
        return about | summary | {"defaults": defaults, "check": check}

    def run(self, form: dict) -> dict:
        frame, labels, _ = self.load(form)
        with self.lock, warnings.catch_warnings(record=True) as caught:  # one model at a time
            warnings.simplefilter("always")
            job = make_job(frame, labels, form)
            result = job.model.fit(job) | {"stats": summary_stats(job)}
        messages = [
            " ".join(str(w.message).split())
            for w in caught
            if not issubclass(w.category, DeprecationWarning | FutureWarning)
        ]
        return result | {"notes": list(dict.fromkeys(job.notes + messages))}


ROUTES = {"/api/files": "files", "/api/profile": "profile", "/api/run": "run"}


class Handler(BaseHTTPRequestHandler):
    server: Portal
    server_version, sys_version = "portal", ""

    def log_message(self, format: str, *args: Any) -> None:  # quiet: errors are printed
        pass

    def do_GET(self) -> None:
        if not self.trusted():
            return
        if self.path != "/":
            return self.send(404, b"Not found", "text/plain")
        nonce = secrets.token_urlsafe(16)
        page = PAGE.read_text(encoding="utf-8").replace("{{nonce}}", nonce)
        csp = (
            f"default-src 'none'; script-src 'nonce-{nonce}'; style-src 'nonce-{nonce}'; "
            "connect-src 'self'; base-uri 'none'; form-action 'none'; frame-ancestors 'none'"
        )
        body = page.replace("{{token}}", self.server.token).encode()
        self.send(200, body, "text/html; charset=utf-8", csp)

    def do_POST(self) -> None:
        if not self.trusted():
            return
        if not secrets.compare_digest(self.headers.get("X-Token", ""), self.server.token):
            return self.send(403, b'{"error": "Missing or wrong token: reload the page"}')
        if self.path not in ROUTES or self.headers.get_content_type() != "application/json":
            return self.send(404, b'{"error": "Not found"}')
        if (length := int(self.headers.get("Content-Length") or 0)) > MAX_BODY:
            return self.send(413, b'{"error": "Request too large"}')
        try:
            request = json.loads(self.rfile.read(length) or b"{}")
            if not isinstance(request, dict):
                raise UserError("The request must be a JSON object")
            status, result = 200, getattr(self.server, ROUTES[self.path])(request)
        except (UserError, ValueError) as exc:  # ValueError covers malformed JSON
            status, result = 400, {"error": str(exc)}
        except Exception as exc:  # shown in the page; the traceback goes to the terminal
            traceback.print_exc()
            status, result = 500, {"error": f"Unexpected error: {type(exc).__name__}: {exc}"}
        self.send(status, json.dumps(jsonable(result), allow_nan=False).encode())

    def trusted(self) -> bool:
        """Only this computer's browser: the Host header must name 127.0.0.1 or localhost."""
        if self.headers.get("Host") in self.server.hosts:
            return True
        self.send(403, b"Forbidden", "text/plain")
        return False

    def send(self, status: int, body: bytes, kind: str = "application/json", csp: str = "") -> None:
        self.send_response(status)
        self.send_header("Content-Type", kind)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        if csp:
            self.send_header("Content-Security-Policy", csp)
        self.end_headers()
        self.wfile.write(body)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    folder = Path(__file__).parent / Path(DATA_FOLDER).expanduser()
    parser.add_argument("--data", type=Path, default=folder, help="folder of .xlsx/.csv files")
    parser.add_argument("--port", type=int, default=8765, help="port on 127.0.0.1; 0 picks one")
    parser.add_argument("--no-browser", action="store_true", help="do not open a browser window")
    args = parser.parse_args(argv)
    folder = args.data.expanduser().resolve()
    if not folder.is_dir():
        raise SystemExit(f"{folder} is not a folder: create it or pass --data")
    server = Portal(folder, args.port)
    url = f"http://127.0.0.1:{server.port}/"
    print(f"Regression portal: {url} (files in {folder}; Ctrl+C stops it)")
    if not args.no_browser:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
